from pathlib import Path
from backend.database import get_db_connection, init_db, get_db_type

DATA_FILE_PATH = Path(__file__).resolve().parent.parent / "data" / "events_seed.xlsx"

FALLBACK_EVENTS_SEED = [
    (1, "Code Sprint 2.0", "Technical", "Fast-paced competitive programming contest. Solve algorithmic problems against the clock.", "2026-11-14", "10:00 AM", "Main Auditorium", 100, 0, "💻"),
    (2, "Robo Race", "Technical", "Design and navigate custom autonomous or RC robots across a challenging obstacle course.", "2026-11-14", "02:00 PM", "Engineering Lab 2", 40, 50, "🤖"),
    (3, "AI Pitch Fest", "Technical", "Present innovative generative AI solutions and startup business models to industry mentors.", "2026-11-15", "11:00 AM", "Seminar Hall A", 60, 0, "💡"),
    (4, "Web Design Rumble", "Technical", "Speed web styling and frontend build-off testing UI/UX creativity, responsiveness, and speed.", "2026-11-14", "04:00 PM", "Computer Lab 1", 50, 0, "🎨"),
    (5, "Battle of Bands", "Cultural", "Live collegiate band face-off featuring rock, indie, and fusion original performances.", "2026-11-15", "06:00 PM", "Open Air Theatre", 300, 0, "🎸"),
    (6, "Dance Fiesta", "Cultural", "Inter-college solo and group dance showcase spanning western, hip-hop, and classical styles.", "2026-11-14", "07:00 PM", "Open Air Theatre", 200, 0, "💃"),
    (7, "Nukkad Natak", "Cultural", "High-energy street play competition bringing pressing social issues alive at the amphitheatre.", "2026-11-15", "03:00 PM", "Amphitheatre", 150, 0, "🎭"),
    (8, "Fashion Fusion", "Cultural", "Runway fashion competition blending traditional ethnic themes with futuristic aesthetics.", "2026-11-15", "08:00 PM", "Main Stage", 250, 100, "👗"),
    (9, "Open Mic Night", "Cultural", "Acoustic music, stand-up comedy, spoken word, and poetry performances under the night sky.", "2026-11-14", "09:00 PM", "Cafeteria Lawn", 100, 0, "🎤"),
    (10, "Five-a-Side Football", "Sports", "Intense quick-match knockout football tournament on the campus sports turf.", "2026-11-14", "08:00 AM", "Sports Ground", 80, 0, "⚽"),
    (11, "Chess Championship", "Sports", "Rapid and blitz chess tournament following official collegiate FIDE rules.", "2026-11-15", "09:00 AM", "Indoor Games Hall", 32, 0, "♟️"),
    (12, "Valorant Showdown", "Gaming", "5v5 tactical FPS LAN esports tournament for ultimate collegiate gaming bragging rights.", "2026-11-14", "01:00 PM", "Gaming Arena", 60, 0, "🎮"),
    (13, "UI/UX Design Workshop", "Workshop", "Hands-on product design session covering wireframing, Figma components, and design systems.", "2026-11-14", "11:00 AM", "Design Lab", 40, 0, "📱"),
    (14, "Intro to Vibe Coding", "Workshop", "Learn to design, build, and deploy full-stack web applications using AI coding tools in hours.", "2026-11-15", "10:00 AM", "Computer Lab 1", 45, 0, "⚡"),
]


def seed_events():
    """
    Seeds the events table from data/events_seed.xlsx or bundled seed data.
    Skips if events table is already populated.
    """
    init_db()

    conn = get_db_connection()
    cursor = conn.cursor()

    # Check if table already has rows
    cursor.execute("SELECT COUNT(*) FROM events")
    row = cursor.fetchone()
    count = row[0] if row else 0
    if count > 0:
        print(f"Events table already contains {count} events. Skipping seeding.")
        conn.close()
        return

    data_rows = []
    if DATA_FILE_PATH.exists():
        try:
            import openpyxl
            print(f"Reading seed data from {DATA_FILE_PATH}...")
            wb = openpyxl.load_workbook(DATA_FILE_PATH)
            if "Events" in wb.sheetnames:
                ws = wb["Events"]
                rows = list(ws.iter_rows(values_only=True))
                if rows and len(rows) > 1:
                    for r in rows[1:]:
                        if r and r[0] is not None:
                            event_id, title, category, description, event_date, event_time, venue, capacity, fee, icon = r[:10]
                            data_rows.append((
                                int(event_id),
                                str(title),
                                str(category),
                                str(description) if description is not None else "",
                                str(event_date),
                                str(event_time),
                                str(venue),
                                int(capacity),
                                int(fee) if fee is not None else 0,
                                str(icon) if icon is not None else "🎉",
                            ))
        except Exception as e:
            print(f"Notice: Could not load {DATA_FILE_PATH} ({e}). Using built-in seed events.")

    if not data_rows:
        print("Using built-in seed events list.")
        data_rows = FALLBACK_EVENTS_SEED

    for row in data_rows:
        cursor.execute(
            """
            INSERT INTO events (id, title, category, description, event_date, event_time, venue, capacity, fee, icon)
            VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
            """,
            row,
        )

    # In PostgreSQL, reset the primary key sequence so subsequent auto-increments start after the seeded IDs
    if get_db_type() == "postgres":
        try:
            cursor.execute("SELECT setval(pg_get_serial_sequence('events', 'id'), COALESCE(max(id), 1)) FROM events;")
        except Exception as e:
            print(f"Notice: Could not reset PostgreSQL sequence: {e}")

    conn.commit()
    cursor.execute("SELECT COUNT(*) FROM events")
    new_count = cursor.fetchone()[0]
    conn.close()
    print(f"Seeding completed successfully! Total events in database: {new_count}")


if __name__ == "__main__":
    import sys
    sys.stdout.reconfigure(encoding="utf-8")
    seed_events()
    conn = get_db_connection()
    cursor = conn.cursor()
    cursor.execute("SELECT id, title, category, event_date, event_time, venue, capacity, fee, icon FROM events")
    rows = cursor.fetchall()
    print("\n--- Current Events in Database ---")
    for r in rows:
        d = dict(r)
        print(f"[{d['id']:2d}] {d['icon']} {d['title']} | {d['category']} | {d['event_date']} {d['event_time']} | {d['venue']} | Cap: {d['capacity']} | Fee: ₹{d['fee']}")
    conn.close()
